#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# Read markdown file
with open('SYSTEM_OVERVIEW.md', 'r', encoding='utf-8') as f:
    content = f.read()

# Create Document
doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Process markdown content
lines = content.split('\n')
i = 0

while i < len(lines):
    line = lines[i]

    # Heading 1 (# Title)
    if line.startswith('# '):
        title = line[2:].strip()
        p = doc.add_heading(title, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1

    # Heading 2 (## Subtitle)
    elif line.startswith('## '):
        heading = line[3:].strip()
        doc.add_heading(heading, level=2)
        i += 1

    # Heading 3 (### SubSubtitle)
    elif line.startswith('### '):
        heading = line[4:].strip()
        doc.add_heading(heading, level=3)
        i += 1

    # Heading 4 (#### SubSubSubtitle)
    elif line.startswith('#### '):
        heading = line[5:].strip()
        doc.add_heading(heading, level=3)
        i += 1

    # Unordered list (- item)
    elif line.startswith('- '):
        items = []
        while i < len(lines) and lines[i].startswith('- '):
            item = lines[i][2:].strip()
            items.append(item)
            i += 1
        for item in items:
            doc.add_paragraph(item, style='List Bullet')

    # Code block (```...```)
    elif line.startswith('```'):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code_lines.append(lines[i])
            i += 1
        if code_lines:
            code_block = '\n'.join(code_lines)
            p = doc.add_paragraph(code_block, style='List Number')
            p.paragraph_format.left_indent = Inches(0.5)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
        i += 1  # skip closing ```

    # Horizontal rule (---)
    elif line.strip() == '---':
        doc.add_paragraph('_' * 80)
        i += 1

    # Bold text (**text**)
    elif '**' in line:
        p = doc.add_paragraph()
        parts = re.split(r'\*\*(.*?)\*\*', line)
        for j, part in enumerate(parts):
            if j % 2 == 1:  # Bold
                run = p.add_run(part)
                run.bold = True
            else:
                p.add_run(part)
        i += 1

    # Regular paragraph
    elif line.strip():
        doc.add_paragraph(line.strip())
        i += 1

    # Empty line
    else:
        i += 1

# Save document
doc.save('JKUAT_Queue_System_Overview.docx')
print("[SUCCESS] Document created: JKUAT_Queue_System_Overview.docx")
